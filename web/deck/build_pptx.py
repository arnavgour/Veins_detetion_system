"""Build the Vena widescreen PowerPoint deck."""

from __future__ import annotations

from io import BytesIO
from pathlib import Path

from lxml import etree
from PIL import Image
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.oxml.ns import qn
from pptx.util import Emu, Inches, Pt

ROOT = Path(__file__).resolve().parents[2]
MEDIA = Path(__file__).resolve().parent / "media"
OUT = ROOT / "Vena.pptx"

W, H = 13.333, 7.5
INK = RGBColor(0x19, 0x15, 0x12)
PAPER = RGBColor(0xF3, 0xEF, 0xE6)
CREAM = RGBColor(0xFF, 0xFB, 0xF4)
SAGE = RGBColor(0x21, 0x5F, 0x4A)
SAGE2 = RGBColor(0x2F, 0x7A, 0x60)
CLAY = RGBColor(0xC4, 0x5C, 0x26)
MUTED = RGBColor(0x74, 0x6B, 0x60)
SOFT = RGBColor(0xDC, 0xEA, 0xE3)
END = RGBColor(0x1F, 0x8A, 0x5B)
BRANCH = RGBColor(0xC4, 0x45, 0x36)
LINE = RGBColor(0xE0, 0xD8, 0xC8)
SERIF = "Georgia"
SANS = "Avenir Next"


def emu(inches: float) -> Emu:
    return Inches(inches)


def set_run_font(run, name, size, color, bold=False, italic=False):
    run.font.name = name
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.italic = italic
    run.font.color.rgb = color
    rPr = run._r.get_or_add_rPr()
    for tag in ("a:latin", "a:ea", "a:cs"):
        el = rPr.find(qn(tag))
        if el is None:
            el = etree.SubElement(rPr, qn(tag))
        el.set("typeface", name)


def set_fill(shape, color, alpha=None):
    shape.fill.solid()
    shape.fill.fore_color.rgb = color
    shape.line.fill.background()
    if alpha is None:
        return
    spPr = shape._element.spPr
    solid = spPr.find(qn("a:solidFill"))
    srgb = solid.find(qn("a:srgbClr"))
    for child in list(srgb):
        if child.tag == qn("a:alpha"):
            srgb.remove(child)
    el = etree.SubElement(srgb, qn("a:alpha"))
    el.set("val", str(int(max(0, min(1, alpha)) * 100000)))


def rect(slide, x, y, w, h, color, alpha=None):
    shape = slide.shapes.add_shape(
        MSO_SHAPE.RECTANGLE, emu(x), emu(y), emu(w), emu(h)
    )
    set_fill(shape, color, alpha)
    return shape


def round_rect(slide, x, y, w, h, color, radius=0.08, alpha=None):
    shape = slide.shapes.add_shape(
        MSO_SHAPE.ROUNDED_RECTANGLE, emu(x), emu(y), emu(w), emu(h)
    )
    try:
        shape.adjustments[0] = radius
    except Exception:
        pass
    set_fill(shape, color, alpha)
    return shape


def textbox(
    slide,
    x,
    y,
    w,
    h,
    text,
    *,
    size=18,
    font=SANS,
    color=INK,
    bold=False,
    italic=False,
    align=PP_ALIGN.LEFT,
    anchor=MSO_ANCHOR.TOP,
    leading=1.05,
):
    box = slide.shapes.add_textbox(emu(x), emu(y), emu(w), emu(h))
    tf = box.text_frame
    tf.word_wrap = True
    tf.auto_size = None
    tf.margin_left = Inches(0.02)
    tf.margin_right = Inches(0.02)
    tf.margin_top = Inches(0)
    tf.margin_bottom = Inches(0)
    tf.anchor = anchor
    lines = text.split("\n")
    for i, line in enumerate(lines):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = align
        p.line_spacing = leading
        p.space_before = Pt(0)
        p.space_after = Pt(0)
        run = p.add_run()
        run.text = line
        set_run_font(run, font, size, color, bold, italic)
    return box


def kicker(slide, x, y, w, label, color=SAGE):
    textbox(
        slide,
        x,
        y,
        w,
        0.32,
        label.upper(),
        size=11,
        font=SANS,
        color=color,
        bold=True,
    )


def footer(slide, n, total, light=False, number_x=11.4):
    color = RGBColor(0xC8, 0xC0, 0xB4) if light else MUTED
    textbox(
        slide,
        0.55,
        7.12,
        5.2,
        0.28,
        "VENA  ·  DORSAL HAND-VEIN STUDIO",
        size=10,
        font=SANS,
        color=color,
        bold=True,
    )
    textbox(
        slide,
        number_x,
        7.12,
        1.4,
        0.28,
        f"{n:02d}  /  {total:02d}",
        size=10,
        font=SANS,
        color=color,
        bold=True,
        align=PP_ALIGN.RIGHT,
    )


def crop_caption(im: Image.Image) -> Image.Image:
    """Drop the dataset caption burned into the bottom of sample photographs."""
    im = im.convert("RGB")
    w, h = im.size
    return im.crop((0, 0, w, int(h * 0.90)))


def picture_cover(slide, path, x, y, w, h):
    dpi = 180
    tw, th = max(2, int(w * dpi)), max(2, int(h * dpi))
    im = crop_caption(Image.open(path))
    iw, ih = im.size
    scale = max(tw / iw, th / ih)
    nw, nh = max(1, int(iw * scale)), max(1, int(ih * scale))
    im = im.resize((nw, nh), Image.Resampling.LANCZOS)
    left = max(0, (nw - tw) // 2)
    top = max(0, (nh - th) // 2)
    im = im.crop((left, top, left + tw, top + th))
    buf = BytesIO()
    im.save(buf, format="JPEG", quality=90)
    buf.seek(0)
    slide.shapes.add_picture(buf, emu(x), emu(y), emu(w), emu(h))


def picture_contain(slide, path, x, y, w, h):
    im = crop_caption(Image.open(path))
    buf = BytesIO()
    im.save(buf, format="PNG")
    buf.seek(0)
    iw, ih = im.size
    ratio = iw / ih
    box_ratio = w / h
    if ratio > box_ratio:
        pw, ph = w, w / ratio
    else:
        ph, pw = h, h * ratio
    px = x + (w - pw) / 2
    py = y + (h - ph) / 2
    slide.shapes.add_picture(buf, emu(px), emu(py), emu(pw), emu(ph))


def blank(prs, color):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    rect(slide, 0, 0, W, H, color)
    return slide


def accent_bar(slide, color=CLAY):
    rect(slide, 0, 0, W, 0.07, color)


def add_stat(slide, x, y, value, label, color=INK):
    textbox(slide, x, y, 2.6, 0.7, value, size=36, font=SERIF, color=color)
    textbox(slide, x, y + 0.68, 2.6, 0.4, label, size=12, font=SANS, color=MUTED, bold=True)


def build():
    prs = Presentation()
    prs.slide_width = Inches(W)
    prs.slide_height = Inches(H)
    total = 15
    orig = MEDIA / "original.jpg"
    hand = MEDIA / "hand.jpg"
    clahe = MEDIA / "clahe.jpg"
    thresh = MEDIA / "threshold.png"
    mask = MEDIA / "mask.png"
    skeleton = MEDIA / "skeleton.png"
    landmarks = MEDIA / "landmarks.png"
    overlay = MEDIA / "overlay.jpg"

    # 01 Title
    s = blank(prs, INK)
    picture_cover(s, overlay, 5.7, 0, 7.633, H)
    rect(s, 0, 0, 6.35, H, INK)
    kicker(s, 0.7, 1.15, 5.0, "Vena  ·  a talk", color=RGBColor(0xE8, 0xA0, 0x6A))
    textbox(
        s, 0.68, 1.7, 5.3, 2.6,
        "The graph\ninside a hand.",
        size=52, font=SERIF, color=CREAM, leading=0.96,
    )
    textbox(
        s, 0.7, 4.5, 4.8, 1.2,
        "A demonstration studio for dorsal vein analysis. Sixteen classical computer-vision stages, each one visible.",
        size=16, font=SANS, color=RGBColor(0xC8, 0xC0, 0xB4), leading=1.25,
    )
    textbox(
        s, 0.7, 6.35, 4.8, 0.4,
        "Not clinical   ·   Not identity   ·   OpenCV",
        size=12, font=SANS, color=RGBColor(0xE8, 0xA0, 0x6A), bold=True,
    )
    footer(s, 1, total, light=True, number_x=4.85)

    # 02 What it is
    s = blank(prs, PAPER)
    accent_bar(s)
    kicker(s, 0.7, 0.55, 6, "Premise")
    textbox(
        s, 0.68, 1.0, 12, 1.4,
        "Watch every process\nland on the hand.",
        size=44, font=SERIF, color=INK, leading=0.98,
    )
    cards = [
        ("A studio", "Upload a dorsal photograph. The same functions as preprocess.py and features.py run live, one frame at a time."),
        ("A pipeline", "Enhance, segment, skeletonize, then mark endpoints and bifurcations. No neural network in the loop."),
        ("A caution", "This is a teaching demonstration — not a clinical device and not an identity-verification product."),
    ]
    for i, (title, body) in enumerate(cards):
        x = 0.7 + i * 4.15
        round_rect(s, x, 3.55, 3.9, 2.85, CREAM, radius=0.08)
        textbox(s, x + 0.28, 3.75, 3.3, 0.28, f"0{i+1}", size=12, font=SANS, color=CLAY, bold=True)
        textbox(s, x + 0.28, 4.1, 3.35, 0.5, title, size=24, font=SERIF, color=INK)
        textbox(s, x + 0.28, 4.7, 3.35, 1.4, body, size=14, font=SANS, color=MUTED, leading=1.25)
    footer(s, 2, total)

    # 03 Rules
    s = blank(prs, INK)
    kicker(s, 0.7, 0.7, 6, "Three rules", color=RGBColor(0xE8, 0xA0, 0x6A))
    textbox(s, 0.68, 1.2, 11, 0.8, "Everything else is implementation.", size=20, font=SANS, color=RGBColor(0xC8, 0xC0, 0xB4))
    rules = [
        ("01", "Dark pixels are veins.", "Binary 0 / intensity < 128. NIR invert exists only so bright captures can join that convention."),
        ("02", "Search the interior of the palm.", "The silhouette is high-contrast bait. A distance-transform margin keeps the fist outline off the graph."),
        ("03", "Believe the overlay.", "A clean skeleton on paper can still be a bracelet. Fit is judged on the photograph."),
    ]
    for i, (num, title, body) in enumerate(rules):
        y = 2.15 + i * 1.45
        textbox(s, 0.7, y, 1.1, 0.5, num, size=18, font=SERIF, color=RGBColor(0xE8, 0xA0, 0x6A))
        textbox(s, 2.0, y, 10.2, 0.5, title, size=28, font=SERIF, color=CREAM)
        textbox(s, 2.0, y + 0.52, 10.0, 0.7, body, size=15, font=SANS, color=RGBColor(0xC8, 0xC0, 0xB4), leading=1.2)
    footer(s, 3, total, light=True)

    # 04 Capture
    s = blank(prs, PAPER)
    accent_bar(s)
    picture_cover(s, orig, 6.35, 0, 6.983, H)
    kicker(s, 0.7, 1.35, 5, "01  ·  Capture")
    textbox(s, 0.68, 1.9, 5.2, 2.0, "The photograph\nis the ground truth.", size=36, font=SERIF, color=INK, leading=0.98)
    textbox(
        s, 0.7, 4.2, 5.0, 1.6,
        "Later overlays only mean anything against this frame. Rings, straps, and blown highlights that look like vessels show up here first.",
        size=16, font=SANS, color=MUTED, leading=1.25,
    )
    textbox(s, 0.7, 6.15, 5.0, 0.4, "Long side capped at 1400 px  ·  otherwise unchanged", size=12, font=SANS, color=SAGE, bold=True)
    footer(s, 4, total, number_x=4.85)

    # 05 The trap
    s = blank(prs, INK)
    picture_cover(s, thresh, 7.1, 0, 6.233, H)
    rect(s, 6.6, 0, 1.0, H, INK, alpha=0.35)
    kicker(s, 0.7, 1.2, 6, "The failure mode", color=RGBColor(0xE8, 0xA0, 0x6A))
    textbox(s, 0.68, 1.8, 6.2, 2.2, "The outline of\nthe hand is not\na vein.", size=40, font=SERIF, color=CREAM, leading=0.98)
    textbox(
        s, 0.7, 4.6, 5.8, 1.8,
        "Adaptive threshold loves the high-contrast edge where the fist meets the background. That contour often never touches the frame, so border cleanup leaves it — and skeletonization traces the palm.",
        size=16, font=SANS, color=RGBColor(0xC8, 0xC0, 0xB4), leading=1.25,
    )
    footer(s, 5, total, light=True, number_x=4.85)

    # 06 Interior ROI
    s = blank(prs, PAPER)
    accent_bar(s, SAGE)
    picture_cover(s, hand, 6.4, 0, 6.933, H)
    kicker(s, 0.7, 1.15, 5.4, "02  ·  Hand interior")
    textbox(s, 0.68, 1.7, 5.4, 1.8, "Veins live a\nmargin inside\nthe palm.", size=36, font=SERIF, color=INK, leading=0.98)
    points = [
        ("Gold band", "Silhouette. Off-limits."),
        ("Green interior", "The only place veins may be scored."),
        ("Distance transform", "Margin ≈ 5.5% of the short side, at least 20 px."),
    ]
    for i, (t, b) in enumerate(points):
        y = 4.05 + i * 0.78
        textbox(s, 0.7, y, 5.2, 0.32, t, size=16, font=SANS, color=SAGE, bold=True)
        textbox(s, 0.7, y + 0.28, 5.2, 0.4, b, size=14, font=SANS, color=MUTED)
    footer(s, 6, total, number_x=4.85)

    # 07 Sixteen stages
    s = blank(prs, PAPER)
    accent_bar(s)
    kicker(s, 0.7, 0.45, 8, "The pipeline")
    textbox(s, 0.68, 0.85, 12, 0.7, "Sixteen frames. Four acts.", size=36, font=SERIF, color=INK)
    acts = [
        ("Capture", SAGE, ["01 Original", "02 Interior ROI"]),
        ("Enhance", SAGE2, ["03 Grayscale / NIR", "04 Median 5×5", "05 CLAHE", "06 Median again"]),
        ("Segment", CLAY, ["07 Adaptive threshold", "08 Morphological smooth", "09 Blob filter", "10 Vein mask"]),
        ("Extract", BRANCH, ["11 Outline isolation", "12 Skeleton prep", "13 Skeletonize", "14 Spur prune", "15 Landmarks", "16 Overlay"]),
    ]
    for i, (title, color, items) in enumerate(acts):
        x = 0.7 + i * 3.15
        rect(s, x, 1.85, 2.95, 0.08, color)
        textbox(s, x, 2.05, 2.95, 0.4, title.upper(), size=12, font=SANS, color=color, bold=True)
        for j, item in enumerate(items):
            textbox(s, x, 2.55 + j * 0.62, 2.95, 0.5, item, size=16, font=SANS, color=INK)
    footer(s, 7, total)

    # 08 Enhance
    s = blank(prs, PAPER)
    accent_bar(s)
    picture_cover(s, clahe, 6.45, 0, 6.883, H)
    kicker(s, 0.7, 1.05, 5.4, "Enhance")
    textbox(s, 0.68, 1.55, 5.4, 1.5, "Lift faint veins\nwithout blowing\nout the skin.", size=34, font=SERIF, color=INK, leading=0.98)
    steps = [
        ("Median 5×5", "Pores vanish. Thin trunks survive."),
        ("CLAHE 8×8, clip 2.0", "Local contrast, limited so grain does not explode."),
        ("Median again", "CLAHE raises leftover speckle; calm it before threshold."),
    ]
    for i, (t, b) in enumerate(steps):
        y = 3.7 + i * 0.85
        textbox(s, 0.7, y, 5.3, 0.32, t, size=16, font=SANS, color=SAGE, bold=True)
        textbox(s, 0.7, y + 0.3, 5.3, 0.45, b, size=14, font=SANS, color=MUTED)
    footer(s, 8, total, number_x=4.85)

    # 09 Segment
    s = blank(prs, PAPER)
    accent_bar(s)
    kicker(s, 0.7, 0.42, 8, "Segment")
    textbox(s, 0.68, 0.82, 12, 0.7, "A local hypothesis, then a cleanup.", size=32, font=SERIF, color=INK)
    round_rect(s, 0.55, 1.75, 6.0, 5.0, CREAM, radius=0.06)
    picture_contain(s, thresh, 0.7, 1.9, 5.7, 4.35)
    textbox(s, 0.75, 6.32, 5.6, 0.3, "Adaptive Gaussian  ·  block 29  ·  C = 2", size=11, font=SANS, color=MUTED, bold=True)
    round_rect(s, 6.8, 1.75, 6.0, 5.0, CREAM, radius=0.06)
    picture_contain(s, mask, 6.95, 1.9, 5.7, 4.35)
    textbox(s, 7.0, 6.32, 5.6, 0.3, "After morphology, blobs, silhouette strip", size=11, font=SANS, color=MUTED, bold=True)
    footer(s, 9, total)

    # 10 1-bit tax
    s = blank(prs, INK)
    kicker(s, 0.7, 0.7, 8, "After threshold", color=RGBColor(0xE8, 0xA0, 0x6A))
    textbox(s, 0.68, 1.2, 12, 1.3, "Binarization is supposed\nto throw information away.", size=36, font=SERIF, color=CREAM, leading=0.98)
    stats = [
        ("1 bit", "from an 8-bit field"),
        ("Speckle", "is a hypothesis, not the answer"),
        ("Trunks", "are recovered by morphology"),
    ]
    for i, (a, b) in enumerate(stats):
        x = 0.7 + i * 4.15
        rect(s, x, 3.3, 3.9, 0.06, CLAY)
        textbox(s, x, 3.55, 3.9, 0.6, a, size=28, font=SERIF, color=CREAM)
        textbox(s, x, 4.25, 3.8, 0.8, b, size=16, font=SANS, color=RGBColor(0xC8, 0xC0, 0xB4), leading=1.2)
    textbox(
        s, 0.7, 5.5, 11.8, 1.0,
        "C stays at 2 so thin branches on a bright palm are not erased. The snowy binary is expected. Later stages throw most of those islands away.",
        size=16, font=SANS, color=RGBColor(0xC8, 0xC0, 0xB4), leading=1.25,
    )
    footer(s, 10, total, light=True)

    # 11 Skeleton
    s = blank(prs, PAPER)
    accent_bar(s)
    round_rect(s, 6.35, 0.45, 6.5, 6.5, CREAM, radius=0.06)
    picture_contain(s, skeleton, 6.5, 0.6, 6.2, 6.2)
    kicker(s, 0.7, 1.1, 5.4, "Extract  ·  skeleton")
    textbox(s, 0.68, 1.65, 5.4, 1.8, "A thick mask is not\na graph. A centerline is.", size=32, font=SERIF, color=INK, leading=0.98)
    bits = [
        ("Prep", "Round vessel walls so jaggies cannot become fur."),
        ("Thin", "scikit-image skeletonize — the medial axis."),
        ("Prune", "Walk 16–36 px from each dead-end. If you hit a fork, erase the twig."),
    ]
    for i, (t, b) in enumerate(bits):
        y = 3.85 + i * 0.9
        textbox(s, 0.7, y, 5.3, 0.32, t, size=16, font=SANS, color=SAGE, bold=True)
        textbox(s, 0.7, y + 0.3, 5.3, 0.5, b, size=14, font=SANS, color=MUTED, leading=1.2)
    footer(s, 11, total)

    # 12 Landmarks
    s = blank(prs, PAPER)
    accent_bar(s)
    round_rect(s, 6.35, 0.45, 6.5, 6.5, CREAM, radius=0.06)
    picture_contain(s, landmarks, 6.5, 0.6, 6.2, 6.2)
    kicker(s, 0.7, 0.85, 5.4, "Extract  ·  landmarks")
    textbox(s, 0.68, 1.35, 5.4, 1.5, "The signature\nis a graph.", size=36, font=SERIF, color=INK, leading=0.98)
    textbox(s, 0.7, 3.05, 5.2, 0.4, "One neighbor", size=14, font=SANS, color=END, bold=True)
    textbox(s, 0.7, 3.35, 5.2, 0.4, "Endpoint  ·  green", size=18, font=SERIF, color=INK)
    textbox(s, 0.7, 3.95, 5.2, 0.4, "Three or more neighbors", size=14, font=SANS, color=BRANCH, bold=True)
    textbox(s, 0.7, 4.25, 5.2, 0.4, "Bifurcation  ·  red", size=18, font=SERIF, color=INK)
    textbox(s, 0.7, 4.95, 5.2, 0.7, "Nearby hits merge at 8 px so one anatomical fork is not a cluster of dots.", size=14, font=SANS, color=MUTED, leading=1.2)
    add_stat(s, 0.7, 5.75, "79", "Endpoints on s4")
    add_stat(s, 3.3, 5.75, "59", "Bifurcations")
    footer(s, 12, total)

    # 13 Overlay
    s = blank(prs, INK)
    picture_cover(s, overlay, 0, 0, W, H)
    rect(s, 0, 5.15, W, 2.35, INK, alpha=0.72)
    kicker(s, 0.7, 5.35, 8, "16  ·  Overlay", color=RGBColor(0xE8, 0xA0, 0x6A))
    textbox(s, 0.68, 5.7, 9.5, 0.7, "The sanity check is the photograph.", size=28, font=SERIF, color=CREAM)
    textbox(
        s, 0.7, 6.4, 11.5, 0.5,
        "Clay mask on vein pixels  ·  Green ends  ·  Red forks  ·  Studio can wipe against the original",
        size=14, font=SANS, color=RGBColor(0xC8, 0xC0, 0xB4), bold=True,
    )
    footer(s, 13, total, light=True)

    # 14 Architecture
    s = blank(prs, PAPER)
    accent_bar(s)
    kicker(s, 0.7, 0.5, 8, "How it is wired")
    textbox(s, 0.68, 0.95, 12, 0.8, "A stream, not a spinner.", size=36, font=SERIF, color=INK)
    flow = [
        ("Browser", "POST the image"),
        ("FastAPI", "NDJSON events"),
        ("pipeline.py", "One frame per stage"),
        ("Studio", "Replay, compare, guide"),
    ]
    for i, (t, b) in enumerate(flow):
        x = 0.7 + i * 3.15
        round_rect(s, x, 2.1, 2.95, 1.7, CREAM, radius=0.1)
        textbox(s, x + 0.22, 2.25, 2.5, 0.28, f"0{i+1}", size=11, font=SANS, color=CLAY, bold=True)
        textbox(s, x + 0.22, 2.55, 2.5, 0.45, t, size=20, font=SERIF, color=INK)
        textbox(s, x + 0.22, 3.1, 2.5, 0.45, b, size=13, font=SANS, color=MUTED)
        if i < 3:
            textbox(s, x + 2.7, 2.65, 0.5, 0.4, "→", size=20, font=SANS, color=CLAY)

    files = [
        ("preprocess.py", "Gray, CLAHE, threshold, silhouette strip"),
        ("features.py", "Skeleton, spur prune, endpoints, forks"),
        ("inversionPreProcess.py", "Bright NIR veins flipped to dark"),
        ("web/", "Studio + About, light-mode"),
    ]
    for i, (t, b) in enumerate(files):
        x = 0.7 + (i % 2) * 6.2
        y = 4.2 + (i // 2) * 1.15
        textbox(s, x, y, 5.8, 0.35, t, size=16, font=SANS, color=SAGE, bold=True)
        textbox(s, x, y + 0.35, 5.8, 0.4, b, size=14, font=SANS, color=MUTED)
    footer(s, 14, total)

    # 15 Close
    s = blank(prs, SAGE)
    kicker(s, 0.7, 1.5, 8, "Open the studio", color=RGBColor(0xE8, 0xA0, 0x6A))
    textbox(s, 0.68, 2.1, 12, 1.8, "Watch every process\nland on the hand.", size=48, font=SERIF, color=CREAM, leading=0.96)
    textbox(s, 0.7, 4.3, 10, 0.5, "http://127.0.0.1:8000", size=22, font=SANS, color=CREAM, bold=True)
    textbox(
        s, 0.7, 5.05, 10, 0.8,
        "Upload or pick a sample. Step with the arrows.\nOn the overlay, compare with the original.",
        size=16, font=SANS, color=RGBColor(0xDC, 0xEA, 0xE3), leading=1.25,
    )
    textbox(s, 0.7, 6.35, 10, 0.4, "Vena  ·  a demonstration, not a diagnosis", size=13, font=SANS, color=RGBColor(0xE8, 0xA0, 0x6A), bold=True)
    footer(s, 15, total, light=True)

    prs.save(OUT)
    print(f"Wrote {OUT} ({OUT.stat().st_size // 1024} KB, {len(prs.slides)} slides)")


if __name__ == "__main__":
    build()
